import docx
import pytest
from docx.shared import Inches

from doc_alttext.package import Package
from doc_alttext.word import scan
from helpers import png, read_part, replace_part


def build(path, setup):
    d = docx.Document()
    setup(d)
    d.save(path)
    return path


def first_doc_pr(d):
    return d.inline_shapes[0]._inline.docPr


def test_body_image_with_heading_caption_and_text(tmp_path):
    def setup(d):
        d.add_heading("Results", level=1)
        d.add_paragraph("Sales grew strongly between 2020 and 2024.")
        d.add_picture(png(), width=Inches(2))
        d.add_paragraph("Figure 1: Sales per year", style="Caption")
        d.add_paragraph("The next section explains the reasons.")

    path = build(tmp_path / "a.docx", setup)
    with Package(path) as pkg:
        result = scan(pkg)
    assert len(result.images) == 1
    img = result.images[0]
    assert img.part == "word/document.xml"
    assert img.media == "word/media/image1.png"
    assert img.problem is None
    assert img.location == "body, paragraph 3"
    assert img.old_alt == "" and img.decorative is False
    assert [t.tag.rsplit("}", 1)[1] for t in img.targets] == ["docPr", "cNvPr"]
    assert "Nearest heading: Results" in img.context
    assert "Caption: Figure 1: Sales per year" in img.context
    assert "Text before: Results Sales grew strongly" in img.context
    assert "Text after: The next section" in img.context
    assert "Figure 1" not in img.context.split("Text after:")[1]
    assert "Sales grew strongly" in result.text


def test_caption_detected_by_seq_field(tmp_path):
    def setup(d):
        d.add_picture(png())
        p = d.add_paragraph("Abbildung ")
        run = p.add_run()
        from docx.oxml import parse_xml

        run._r.append(parse_xml('<w:instrText xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"> SEQ Abbildung \\* ARABIC </w:instrText>'))
        p.add_run(": Umsatz")

    with Package(build(tmp_path / "a.docx", setup)) as pkg:
        img = scan(pkg).images[0]
    assert "Caption: Abbildung : Umsatz" in img.context


def test_existing_alt_and_decorative_are_read(tmp_path):
    def setup(d):
        d.add_picture(png())
        d.add_picture(png(color=(0, 0, 255)))
        d.inline_shapes[0]._inline.docPr.set("descr", "A red square")
        d.inline_shapes[1]._inline.docPr.set("descr", "image2.png")

    with Package(build(tmp_path / "a.docx", setup)) as pkg:
        images = scan(pkg).images
    assert [i.old_alt for i in images] == ["A red square", ""]


def test_image_in_table_and_header(tmp_path):
    def setup(d):
        cell = d.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].add_run().add_picture(png())
        header = d.sections[0].header
        header.paragraphs[0].text = "Company report"
        header.paragraphs[0].add_run().add_picture(png(size=(80, 40)))

    with Package(build(tmp_path / "a.docx", setup)) as pkg:
        images = scan(pkg).images
    locations = sorted(i.location for i in images)
    assert locations[0].startswith("body, table, paragraph 1")
    assert locations[1].startswith("header (header1)")
    header_img = next(i for i in images if i.location.startswith("header"))
    assert "page header" in header_img.context
    assert "repeat on every page" in header_img.context
    assert "Company report" in header_img.context
    assert header_img.media.startswith("word/media/")


def test_linked_image_is_a_problem(tmp_path):
    path = build(tmp_path / "a.docx", lambda d: d.add_picture(png()))
    xml = read_part(path, "word/document.xml").replace("r:embed=", "r:link=")
    replace_part(path, "word/document.xml", xml.encode())
    with Package(path) as pkg:
        img = scan(pkg).images[0]
    assert img.media is None
    assert "linked" in img.problem


def test_vml_image_is_reported(tmp_path):
    path = build(tmp_path / "a.docx", lambda d: d.add_paragraph("x"))
    vml = (
        '<w:p><w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" id="_x0000_i1025" alt="old">'
        '<v:imagedata r:id="rId99"/></v:shape></w:pict></w:r></w:p>'
    )
    xml = read_part(path, "word/document.xml").replace("<w:sectPr", vml + "<w:sectPr", 1)
    replace_part(path, "word/document.xml", xml.encode())
    with Package(path) as pkg:
        img = scan(pkg).images[0]
    assert "VML" in img.problem
    assert img.old_alt == "old"
    assert img.targets == []


def test_group_pictures_get_own_targets(tmp_path):
    path = build(tmp_path / "a.docx", lambda d: d.add_picture(png()))
    xml = read_part(path, "word/document.xml")
    start = xml.index("<pic:pic")
    end = xml.index("</pic:pic>") + len("</pic:pic>")
    pic = xml[start:end]
    second = pic.replace('name="', 'name="second ', 1)
    xml = xml[:start] + pic + second + xml[end:]
    replace_part(path, "word/document.xml", xml.encode())
    with Package(path) as pkg:
        images = scan(pkg).images
    assert len(images) == 2
    assert all(len(i.targets) == 1 and i.targets[0].tag.endswith("cNvPr") for i in images)
    assert images[0].location.endswith(", group")


def test_fallback_copies_are_ignored(tmp_path):
    path = build(tmp_path / "a.docx", lambda d: d.add_picture(png()))
    xml = read_part(path, "word/document.xml")
    start = xml.index("<w:drawing>")
    end = xml.index("</w:drawing>") + len("</w:drawing>")
    drawing = xml[start:end]
    alternate = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        f"<mc:Choice Requires=\"wps\">{drawing}</mc:Choice><mc:Fallback>{drawing}</mc:Fallback></mc:AlternateContent>"
    )
    replace_part(path, "word/document.xml", (xml[:start] + alternate + xml[end:]).encode())
    with Package(path) as pkg:
        assert len(scan(pkg).images) == 1


def test_lang_attributes_counted(tmp_path):
    path = build(tmp_path / "a.docx", lambda d: d.add_paragraph("Hallo"))
    xml = read_part(path, "word/document.xml").replace("<w:t>Hallo", '<w:rPr><w:lang w:val="de-DE"/></w:rPr><w:t>Hallo')
    replace_part(path, "word/document.xml", xml.encode())
    with Package(path) as pkg:
        assert scan(pkg).langs["de-DE"] == 1


WPS_DRAWING = (
    '<w:r><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" '
    'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    '<wp:extent cx="914400" cy="914400"/><wp:wrapNone/><wp:docPr id="50" name="{name}"/>'
    '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
    '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:cNvSpPr/>'
    '<wps:spPr><a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom></wps:spPr>{text}<wps:bodyPr/></wps:wsp>'
    '</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'
)


def test_word_shape_without_text(tmp_path):
    path = build(tmp_path / "a.docx", lambda d: d.add_paragraph("Before the circle"))
    empty = WPS_DRAWING.replace("{name}", "Oval 1").replace("{text}", "")
    with_text = WPS_DRAWING.replace("{name}", "Box 2").replace(
        "{text}", "<wps:txbx><w:txbxContent><w:p><w:r><w:t>Label</w:t></w:r></w:p></w:txbxContent></wps:txbx>"
    )
    xml = read_part(path, "word/document.xml").replace("<w:t>Before the circle</w:t></w:r>",
                                                       "<w:t>Before the circle</w:t></w:r>" + empty + with_text, 1)
    replace_part(path, "word/document.xml", xml.encode())
    with Package(path) as pkg:
        images = scan(pkg).images
    assert [(i.shape, i.plain_shape, i.location) for i in images] == [("Oval 1", True, "body, paragraph 1, shape")]
    assert images[0].targets[0].tag.endswith("docPr")
