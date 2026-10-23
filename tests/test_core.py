import json

import docx
import pptx
import pytest

from doc_alttext import AltTexter, add_alt_text, default_output
from doc_alttext.errors import AltTextError, LimitError, UnsupportedFileError
from doc_alttext.package import DECORATIVE_URI
from helpers import parts, png, read_part, replace_part


def make_docx(path, pictures=1, text="Sales grew strongly in 2024.", same=False):
    d = docx.Document()
    if text:
        d.add_heading("Results", level=1)
        d.add_paragraph(text)
    for i in range(pictures):
        d.add_picture(png(color=(200, 30, 30) if same else (i * 40 % 256, 30, 30)))
    d.save(path)
    return path


def make_pptx(path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Team"
    slide.shapes.add_picture(png(), 0, 0)
    prs.save(path)
    return path


def run(server, path, **kwargs):
    kwargs.setdefault("max_retries", 0)
    return add_alt_text(path, base_url=server.base_url, model="vision", **kwargs)


def test_docx_end_to_end(server, tmp_path):
    src = make_docx(tmp_path / "report.docx")
    before = parts(src)
    result = run(server, src)

    assert result.output == tmp_path / "report.alt.docx"
    assert result.language == "German"
    assert result.ok
    [item] = result.items
    assert (item.status, item.new_alt, item.chars) == ("added", "A test image", 12)
    assert item.location == "body, paragraph 3"
    assert item.image == "word/media/image1.png"

    xml = read_part(result.output, "word/document.xml")
    assert xml.count('descr="A test image"') == 2  # wp:docPr and pic:cNvPr
    assert parts(src) == before  # input untouched
    docx.Document(str(result.output))  # still opens

    image_request = server.image_requests[0]["body"]
    assert "Write it in German." in image_request["messages"][0]["content"]
    assert "Nearest heading: Results" in image_request["messages"][1]["content"][0]["text"]


def test_other_parts_are_copied_byte_for_byte(server, tmp_path):
    src = make_docx(tmp_path / "a.docx")
    result = run(server, src)
    before, after = parts(src), parts(result.output)
    assert list(before) == list(after)
    assert [n for n in before if before[n] != after[n]] == ["word/document.xml"]


def test_pptx_decorative(server, tmp_path):
    server.reply = lambda body: "English" if isinstance(body["messages"][-1]["content"], str) else '{"decorative": true}'
    result = run(server, make_pptx(tmp_path / "deck.pptx"))
    assert result.items[0].status == "decorative"
    xml = read_part(result.output, "ppt/slides/slide1.xml")
    assert DECORATIVE_URI in xml and "adec:decorative" in xml
    pptx.Presentation(str(result.output))


def test_existing_alt_kept_and_overwrite(server, tmp_path):
    src = make_docx(tmp_path / "a.docx")
    xml = read_part(src, "word/document.xml").replace('<wp:docPr id="1" name="Picture 1"', '<wp:docPr id="1" name="Picture 1" descr="Old text"', 1)
    replace_part(src, "word/document.xml", xml.encode())

    result = run(server, src, dry_run=True)
    assert result.items[0].status == "kept"
    assert server.requests == []  # nothing to ask, no language detection either

    result = run(server, src, overwrite=True, dry_run=True)
    assert result.items[0].status == "added"
    assert result.items[0].old_alt == "Old text"


def test_dry_run_writes_nothing(server, tmp_path):
    src = make_docx(tmp_path / "a.docx")
    result = run(server, src, dry_run=True)
    assert result.output is None
    assert result.items[0].new_alt == "A test image"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["a.docx"]


def test_in_place(server, tmp_path):
    src = make_docx(tmp_path / "a.docx")
    result = run(server, src, in_place=True)
    assert result.output == src
    assert 'descr="A test image"' in read_part(src, "word/document.xml")
    assert sorted(p.name for p in tmp_path.iterdir()) == ["a.docx"]


def test_output_rules(server, tmp_path):
    src = make_docx(tmp_path / "a.docx")
    with pytest.raises(AltTextError, match="must not be the input"):
        run(server, src, output=src)
    with pytest.raises(AltTextError, match="not both"):
        AltTexter(base_url=server.base_url, model="m").process(src, tmp_path / "b.docx", in_place=True)
    with pytest.raises(AltTextError, match="must end in .docx"):
        run(server, src, output=tmp_path / "b.pptx")
    (tmp_path / "a.alt.docx").write_bytes(b"x")
    assert default_output(src) == tmp_path / "a.alt-1.docx"
    explicit = tmp_path / "out" / "named.docx"
    assert run(server, src, output=explicit).output == explicit and explicit.exists()


def test_same_image_is_asked_once(server, tmp_path):
    src = make_docx(tmp_path / "a.docx", pictures=3, same=True)
    result = run(server, src)
    assert [i.status for i in result.items] == ["added"] * 3
    assert len(server.image_requests) == 1
    assert "appears 3 times" in server.image_requests[0]["body"]["messages"][1]["content"][0]["text"]


def test_tiny_image_is_decorative_without_model(server, tmp_path):
    d = docx.Document()
    d.add_picture(png(size=(600, 4)))
    d.save(tmp_path / "a.docx")
    result = run(server, tmp_path / "a.docx")
    assert result.items[0].status == "decorative"
    assert result.items[0].note == "tiny or thin image"
    assert server.requests == []


def test_lang_setting_skips_detection(server, tmp_path):
    result = run(server, make_docx(tmp_path / "a.docx"), lang="French", dry_run=True)
    assert result.language == "French"
    assert len(server.requests) == len(server.image_requests) == 1
    assert "Write it in French." in server.requests[0]["body"]["messages"][0]["content"]


def test_language_fallbacks(server, tmp_path):
    src = make_docx(tmp_path / "a.docx", text="")
    assert run(server, src, dry_run=True).language == "English"
    xml = read_part(src, "word/document.xml").replace("<w:body>", '<w:body><w:p><w:r><w:rPr><w:lang w:val="de-DE"/></w:rPr></w:r></w:p>', 1)
    replace_part(src, "word/document.xml", xml.encode())
    assert run(server, src, dry_run=True).language == "de-DE"


def test_model_error_is_reported_and_file_still_written(server, tmp_path):
    server.status = 500
    result = run(server, make_docx(tmp_path / "a.docx"), lang="English")
    assert not result.ok
    assert result.items[0].status == "error"
    assert "failed" in result.items[0].note
    assert result.output.exists()


def test_too_long_is_flagged(server, tmp_path):
    long_alt = "x" * 200
    server.reply = lambda body: json.dumps({"decorative": False, "alt": long_alt})
    result = run(server, make_docx(tmp_path / "a.docx"), lang="English", dry_run=True)
    assert result.items[0].note == "too_long"
    assert result.items[0].new_alt == long_alt


def test_unsupported_image_format(server, tmp_path):
    src = make_docx(tmp_path / "a.docx")
    with __import__("zipfile").ZipFile(src) as z:
        items = [(i, z.read(i)) for i in z.infolist()]
    import zipfile

    with zipfile.ZipFile(src, "w") as z:
        for info, data in items:
            if info.filename == "word/media/image1.png":
                info.filename = "word/media/image1.emf"
            elif info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"media/image1.png", b"media/image1.emf")
            z.writestr(info, data)
    result = run(server, src, dry_run=True)
    assert result.items[0].status == "unsupported"
    assert ".emf" in result.items[0].note


def test_limits_and_file_types(server, tmp_path):
    src = make_docx(tmp_path / "a.docx", pictures=2)
    with pytest.raises(LimitError, match="--max-images"):
        run(server, src, max_images=1)
    with pytest.raises(LimitError, match="--max-file-mb"):
        run(server, src, max_file_mb=0.001)
    other = tmp_path / "a.pdf"
    other.write_bytes(b"%PDF")
    with pytest.raises(UnsupportedFileError, match="Unsupported file type .pdf"):
        run(server, other)
    with pytest.raises(FileNotFoundError):
        run(server, tmp_path / "missing.docx")


def test_apply_alt_texts(server, tmp_path):
    from doc_alttext import apply_alt_texts
    from doc_alttext.package import Package
    from doc_alttext.word import scan

    src = make_docx(tmp_path / "a.docx", pictures=3)
    result = run(server, src, dry_run=True)
    assert len(result.items) == 3
    asked = len(server.requests)
    out = apply_alt_texts(src, [{"alt": "  Edited   text "}, {"decorative": True}, None])
    assert out == tmp_path / "a.alt.docx"
    with Package(out) as pkg:
        images = scan(pkg).images
    assert images[0].old_alt == "Edited text"
    assert images[1].decorative and images[1].old_alt == ""
    assert images[2].old_alt == "" and not images[2].decorative
    assert len(server.requests) == asked  # apply asks no model


def test_apply_alt_texts_checks_count_and_output(tmp_path):
    from doc_alttext import apply_alt_texts

    src = make_docx(tmp_path / "a.docx", pictures=1)
    with pytest.raises(AltTextError, match="1 images, but 2 decisions"):
        apply_alt_texts(src, [None, None])
    with pytest.raises(AltTextError, match="must not be the input"):
        apply_alt_texts(src, [None], src)
    assert apply_alt_texts(src, [{"alt": "x"}], in_place=True) == src
    assert 'descr="x"' in read_part(src, "word/document.xml")


def test_type_comes_from_content_not_name(server, tmp_path):
    deck = make_pptx(tmp_path / "deck.pptx")
    renamed = tmp_path / "deck.docx"
    renamed.write_bytes(deck.read_bytes())
    result = run(server, renamed, dry_run=True)
    assert [i.location for i in result.items] == ["slide 1"]

    other = tmp_path / "sheet.docx"
    import zipfile

    with zipfile.ZipFile(other, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="r1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    with pytest.raises(UnsupportedFileError, match="neither a Word nor a PowerPoint"):
        run(server, other, dry_run=True)


def test_plain_shapes_become_decorative_without_model(server, tmp_path):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for _ in range(3):
        slide.shapes.add_shape(MSO_SHAPE.OVAL, 0, 0, Inches(1), Inches(1))
    kept = slide.shapes.add_shape(MSO_SHAPE.OVAL, 0, 0, Inches(1), Inches(1))
    kept._element.nvSpPr.cNvPr.set("descr", "Red status light")
    prs.save(tmp_path / "a.pptx")

    result = run(server, tmp_path / "a.pptx", max_images=1)  # shapes do not count against the limit
    assert [i.status for i in result.items] == ["decorative"] * 3 + ["kept"]
    assert result.items[0].note == "shape without text"
    assert server.requests == []
    xml = read_part(result.output, "ppt/slides/slide1.xml")
    assert xml.count(DECORATIVE_URI) == 3
    assert 'descr="Red status light"' in xml


def test_same_image_on_two_slides_gets_its_own_context(server, tmp_path):
    prs = pptx.Presentation()
    for title in ("Goals", "Scope"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        slide.shapes.add_picture(png(), 0, 0)
    prs.save(tmp_path / "a.pptx")
    run(server, tmp_path / "a.pptx", lang="English", dry_run=True)
    texts = [r["body"]["messages"][1]["content"][0]["text"] for r in server.image_requests]
    assert len(texts) == 2  # requests run in parallel, so their order is not fixed
    assert sum("Slide title: Goals" in t for t in texts) == 1
    assert sum("Slide title: Scope" in t for t in texts) == 1
    assert all("appears 2 times" in t for t in texts)
