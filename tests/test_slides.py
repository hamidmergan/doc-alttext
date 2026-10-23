import pptx
from pptx.util import Inches

from doc_alttext.package import Package
from doc_alttext.slides import scan
from helpers import png, read_part, replace_part


def deck(path, *slides):
    prs = pptx.Presentation()
    for setup in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        setup(slide)
    prs.save(path)
    return path


def test_picture_with_title_text_and_notes(tmp_path):
    def first(slide):
        slide.shapes.title.text = "Intro"

    def second(slide):
        slide.shapes.title.text = "Sales 2020-2024"
        slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1)).text_frame.text = "Sales doubled"
        slide.shapes.add_picture(png(), Inches(1), Inches(2))
        slide.notes_slide.notes_text_frame.text = "Mention the new markets."

    path = deck(tmp_path / "a.pptx", first, second)
    with Package(path) as pkg:
        result = scan(pkg)
    assert len(result.images) == 1
    img = result.images[0]
    assert img.location == "slide 2"
    assert img.part == "ppt/slides/slide2.xml"
    assert img.media == "ppt/media/image1.png"
    assert img.shape.startswith("Picture")
    assert "slide 2 of a presentation" in img.context
    assert "Slide title: Sales 2020-2024" in img.context
    assert "Other text on the slide: Sales doubled" in img.context
    assert "Speaker notes: Mention the new markets." in img.context
    assert "Intro" in result.text and "Sales doubled" in result.text


def test_existing_alt_kept(tmp_path):
    def setup(slide):
        pic = slide.shapes.add_picture(png(), 0, 0)
        pic._element.nvPicPr.cNvPr.set("descr", "Team photo")

    with Package(deck(tmp_path / "a.pptx", setup)) as pkg:
        assert scan(pkg).images[0].old_alt == "Team photo"


def test_group_and_picture_fill(tmp_path):
    def setup(slide):
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(png(), 0, 0)

    path = deck(tmp_path / "a.pptx", setup)
    with Package(path) as pkg:
        img = scan(pkg).images[0]
    assert img.location == "slide 1, group"

    # a rectangle filled with the same picture
    xml = read_part(path, "ppt/slides/slide1.xml")
    rid = xml.split('r:embed="', 1)[1].split('"', 1)[0]
    shape = (
        '<p:sp><p:nvSpPr><p:cNvPr id="99" name="Filled"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill></p:spPr></p:sp>'
    )
    xml = xml.replace("</p:spTree>", shape + "</p:spTree>")
    replace_part(path, "ppt/slides/slide1.xml", xml.encode())
    with Package(path) as pkg:
        images = scan(pkg).images
    filled = next(i for i in images if i.shape == "Filled")
    assert filled.location == "slide 1, picture fill"
    assert filled.media == "ppt/media/image1.png"


def test_slide_order_follows_presentation(tmp_path):
    def a(slide):
        slide.shapes.title.text = "A"

    def b(slide):
        slide.shapes.add_picture(png(), 0, 0)

    path = deck(tmp_path / "a.pptx", a, b)
    # move slide 2 to the front
    xml = read_part(path, "ppt/presentation.xml")
    first, second = xml.split("<p:sldId ")[1:3]
    first, second = "<p:sldId " + first.split("/>")[0] + "/>", "<p:sldId " + second.split("/>")[0] + "/>"
    xml = xml.replace(first + second, second + first)
    replace_part(path, "ppt/presentation.xml", xml.encode())
    with Package(path) as pkg:
        assert scan(pkg).images[0].location == "slide 1"


def test_video_is_skipped_and_lang_counted(tmp_path):
    def setup(slide):
        slide.shapes.title.text = "Hallo"
        slide.shapes.add_picture(png(), 0, 0)

    path = deck(tmp_path / "a.pptx", setup)
    xml = read_part(path, "ppt/slides/slide1.xml")
    xml = xml.replace("<p:nvPr/></p:nvPicPr>", '<p:nvPr><a:videoFile r:link="rId9"/></p:nvPr></p:nvPicPr>')
    xml = xml.replace("<a:t>Hallo", '<a:rPr lang="de-DE"/><a:t>Hallo')
    replace_part(path, "ppt/slides/slide1.xml", xml.encode())
    with Package(path) as pkg:
        result = scan(pkg)
    assert result.images == []
    assert result.langs["de-DE"] == 1


def test_plain_shapes_without_text(tmp_path):
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE

    def setup(slide):
        slide.shapes.title.text = "Timeline"
        slide.shapes.add_shape(MSO_SHAPE.OVAL, 0, 0, Inches(1), Inches(1))  # no text: decorative candidate
        slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0, 0, Inches(2), Inches(1)).text_frame.text = "Budget"
        slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Inches(3), 0)
        slide.shapes.add_textbox(0, 0, Inches(1), Inches(1))  # empty text box: left out

    with Package(deck(tmp_path / "a.pptx", setup)) as pkg:
        images = scan(pkg).images
    assert [(i.location, i.plain_shape, i.media) for i in images] == [
        ("slide 1, shape", True, None),
        ("slide 1, shape", True, None),
    ]
    assert images[0].shape.startswith("Oval")
    assert "Connector" in images[1].shape
