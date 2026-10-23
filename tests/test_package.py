import docx
import pytest
from lxml import etree

from doc_alttext.errors import UnsupportedFileError
from doc_alttext.package import (
    A,
    DECORATIVE_URI,
    Package,
    existing_alt,
    is_decorative,
    resolve_target,
    set_alt,
)
from helpers import parts, png

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Taken from PowerPoint's own theme file Bjorn.thmx.
POWERPOINT_DECORATIVE = (
    f'<p:cNvPr xmlns:p="{P_NS}" xmlns:a="{A_NS}" id="7" name="Straight Connector 6"><a:extLst>'
    '<a:ext uri="{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}"><a16:creationId '
    'xmlns:a16="http://schemas.microsoft.com/office/drawing/2014/main" id="{FD144482-F463-E942-01AE-69DF54F4171C}"/>'
    '</a:ext><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}"><adec:decorative '
    'xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/></a:ext></a:extLst></p:cNvPr>'
)


def cnvpr(attrs="", inner=""):
    return etree.fromstring(f'<p:cNvPr xmlns:p="{P_NS}" xmlns:a="{A_NS}" id="4" name="Picture 3" {attrs}>{inner}</p:cNvPr>')


def test_reads_powerpoint_decorative_markup():
    assert is_decorative(etree.fromstring(POWERPOINT_DECORATIVE))
    assert not is_decorative(cnvpr())


def test_set_decorative_matches_powerpoint_markup():
    el = cnvpr('descr="old"')
    set_alt([el], "", decorative=True)
    xml = etree.tostring(el).decode()
    assert "descr" not in xml
    assert (
        f'<a:extLst><a:ext uri="{DECORATIVE_URI}"><adec:decorative '
        'xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/></a:ext></a:extLst>'
    ) in xml


def test_set_decorative_keeps_other_extensions_and_hyperlink_order():
    el = etree.fromstring(POWERPOINT_DECORATIVE.replace('<a:extLst>', '<a:hlinkClick xmlns:r="r" r:id="rId1"/><a:extLst>'))
    set_alt([el], "", decorative=True)  # already decorative: replaced, not doubled
    assert [etree.QName(c).localname for c in el] == ["hlinkClick", "extLst"]
    uris = [e.get("uri") for e in el.find(A + "extLst")]
    assert uris == ["{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}", DECORATIVE_URI]


def test_set_alt_removes_decorative_flag():
    el = etree.fromstring(POWERPOINT_DECORATIVE)
    set_alt([el], "Bar chart: sales doubled", decorative=False)
    assert el.get("descr") == "Bar chart: sales doubled"
    assert not is_decorative(el)
    assert [e.get("uri") for e in el.find(A + "extLst")] == ["{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}"]


def test_set_alt_drops_empty_ext_list():
    el = cnvpr(inner=f'<a:extLst><a:ext uri="{DECORATIVE_URI}"/></a:extLst>')
    set_alt([el], "Text", decorative=False)
    assert el.find(A + "extLst") is None


def test_set_alt_on_several_targets():
    a, b = cnvpr(), cnvpr()
    set_alt([a, b], "Same", decorative=False)
    assert a.get("descr") == b.get("descr") == "Same"


@pytest.mark.parametrize(
    "descr, expected",
    [
        ("A red apple", "A red apple"),
        ("", ""),
        ("   ", ""),
        ("image1.png", ""),
        ("IMG_2034.JPG", ""),
        ("Ein Bild, das Text enthält.\n\nAutomatisch generierte Beschreibung", ""),
        ("A picture containing text\n\nDescription automatically generated", ""),
        ("Chart of results.png in 2024", "Chart of results.png in 2024"),
    ],
)
def test_existing_alt(descr, expected):
    el = cnvpr()
    el.set("descr", descr)
    assert existing_alt(el) == expected


def test_resolve_target():
    assert resolve_target("word", "media/image1.png") == "word/media/image1.png"
    assert resolve_target("ppt/slides", "../media/image1.png") == "ppt/media/image1.png"
    assert resolve_target("ppt/slides", "/ppt/media/x.png") == "ppt/media/x.png"


def test_package_rels_and_save_keeps_other_parts(tmp_path):
    src = tmp_path / "a.docx"
    d = docx.Document()
    d.add_paragraph("Hello")
    d.add_picture(png())
    d.save(src)

    with Package(src) as pkg:
        main = pkg.main_part()
        assert main == "word/document.xml"
        targets = [t for _type, t, _ext in pkg.rels(main).values()]
        assert "word/media/image1.png" in targets
        pkg.xml(main).set("marker", "1")
        pkg.changed.add(main)
        pkg.save(tmp_path / "b.docx")

    before, after = parts(src), parts(tmp_path / "b.docx")
    assert list(before) == list(after)
    for name in before:
        if name != "word/document.xml":
            assert before[name] == after[name], name
    assert b'marker="1"' in after["word/document.xml"]


def test_not_a_zip(tmp_path):
    bad = tmp_path / "x.docx"
    bad.write_bytes(b"no zip")
    with pytest.raises(UnsupportedFileError, match="not a valid"):
        Package(bad)
